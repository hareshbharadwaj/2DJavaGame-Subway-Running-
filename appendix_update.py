# -*- coding: utf-8 -*-
from docx import Document

path = r"C:\Users\hesh2\OneDrive\Desktop\subway game\JAVA REPORT (1)_merged.docx"
doc = Document(path)

heading_a = "APPENDIX A: GitHub Repository and Project Structure"
heading_b = "APPENDIX B: Assets Used and Screenshots"
heading_c = "APPENDIX C: Source Code Snippet (Import Section Only)"

if not any(heading_a in p.text for p in doc.paragraphs):
    doc.add_paragraph()

    p = doc.add_paragraph()
    r = p.add_run(heading_a)
    r.bold = True
    doc.add_paragraph("GitHub Link: https://github.com/hareshbharadwaj/2DJavaGame-Subway-Running-.git")
    doc.add_paragraph("Project Folder Structure:")
    doc.add_paragraph("2DJavaGame-Subway-Running-/")
    doc.add_paragraph("|-- assets/")
    doc.add_paragraph("|   |-- collectibles/")
    doc.add_paragraph("|   |-- environment/")
    doc.add_paragraph("|   |-- obstacles/")
    doc.add_paragraph("|   |-- player/")
    doc.add_paragraph("|   |-- ui/")
    doc.add_paragraph("|-- sounds/")
    doc.add_paragraph("|-- db.properties")
    doc.add_paragraph("|-- DbManager.java")
    doc.add_paragraph("|-- PlayerInfo.java")
    doc.add_paragraph("|-- SessionManager.java")
    doc.add_paragraph("|-- SimpleRunnerGame.java")
    doc.add_paragraph("|-- SoundManager.java")
    doc.add_paragraph("|-- README.md")
    doc.add_paragraph("|-- player.txt")
    doc.add_paragraph("|-- .git/")

    doc.add_paragraph()

    p = doc.add_paragraph()
    r = p.add_run(heading_b)
    r.bold = True
    doc.add_paragraph("B.1 Assets Used:")
    doc.add_paragraph("- Collectibles: coin1.png to coin8.png, coin (static).png")
    doc.add_paragraph("- Environment: background.png, road (1).png")
    doc.add_paragraph("- Obstacles: barrel (1).png, barrier (1).png, car_blue (1).png, car_red (1).png, car_taxi (1).png, cone (1).png, truck (1).png")
    doc.add_paragraph("- Player: player_run (2).png, player_run (1)_processed.png, player_left (1).png, player_right (1).png, player_jump (1).png")
    doc.add_paragraph("- UI: heart (1).png, magnet (1).png, shield (1).png")
    doc.add_paragraph("- Audio: background music and sound effects")
    doc.add_paragraph("B.2 Screenshots:")
    doc.add_paragraph("Figure B.1: Home Page")
    doc.add_paragraph("[Insert Home Page Screenshot Here]")
    doc.add_paragraph("Figure B.2: Game Page")
    doc.add_paragraph("[Insert Game Page Screenshot Here]")
    doc.add_paragraph("Figure B.3: Leaderboard Page")
    doc.add_paragraph("[Insert Leaderboard Page Screenshot Here]")

    doc.add_paragraph()

    p = doc.add_paragraph()
    r = p.add_run(heading_c)
    r.bold = True
    doc.add_paragraph("import java.awt.*;")
    doc.add_paragraph("import java.awt.event.*;")
    doc.add_paragraph("import java.awt.geom.Rectangle2D;")
    doc.add_paragraph("import java.awt.geom.RoundRectangle2D;")
    doc.add_paragraph("import java.util.ArrayList;")
    doc.add_paragraph("import java.util.Iterator;")
    doc.add_paragraph("import java.util.List;")
    doc.add_paragraph("import java.util.Random;")
    doc.add_paragraph("import javax.swing.*;")
    doc.add_paragraph("import javax.imageio.ImageIO;")
    doc.add_paragraph("import java.io.File;")
    doc.add_paragraph("import java.awt.Image;")

    doc.save(path)
    print("APPENDIX_UPDATED")
else:
    print("APPENDIX_ALREADY_PRESENT")
